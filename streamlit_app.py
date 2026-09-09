import streamlit as st
import pandas as pd
import pyodbc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import calendar
from datetime import datetime
import smtplib
import uuid
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

# ─────────────────────────────────────────────
#  Configuração da Página
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="Base Gestor - Orçado x Realizado",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─────────────────────────────────────────────
#  Constantes
# ─────────────────────────────────────────────
MESES_PT = {1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril", 5: "Maio", 6: "Junho",
            7: "Julho", 8: "Agosto", 9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro"}
MESES_SHORT = {1: "Jan", 2: "Fev", 3: "Mar", 4: "Abr", 5: "Mai", 6: "Jun",
               7: "Jul", 8: "Ago", 9: "Set", 10: "Out", 11: "Nov", 12: "Dez"}

ORC_KEYS = ["ORC_JAN", "ORC_FEV", "ORC_MAR", "ORC_ABR", "ORC_MAI", "ORC_JUN",
            "ORC_JUL", "ORC_AGO", "ORC_SET", "ORC_OUT", "ORC_NOV", "ORC_DEZ"]

# ─────────────────────────────────────────────
#  Funções de Formatação
# ─────────────────────────────────────────────
def _brl(v) -> str:
    try:
        if v is None or pd.isna(v): return "R$ 0,00"
        n = float(v)
        s = f"R$ {abs(n):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"-{s}" if n < 0 else s
    except: return "R$ 0,00"

def _pct(pct: float) -> str:
    try:
        sign = "+" if pct > 0 else ""
        return f"{sign}{pct:,.2f}%".replace(",", "X").replace(".", ",").replace("X", ".")
    except: return "0,00%"

def _date(v) -> str:
    try:
        s = str(v).strip()
        return f"{s[6:8]}/{s[4:6]}/{s[0:4]}" if len(s) == 8 and s.isdigit() else s
    except: return str(v) if v else ""

def _s(v) -> str:
    if v is None or pd.isna(v): return ""
    s = str(v).strip()
    return "" if s in ("None", "nan", "NULL") else s

def _f(v) -> float:
    try: return float(v) if v is not None and not pd.isna(v) else 0.0
    except: return 0.0

# ─────────────────────────────────────────────
#  Conexão com Banco de Dados (ODBC / Protheus)
# ─────────────────────────────────────────────
def get_db_credentials():
    secrets_db = st.secrets.get("database", {}) if hasattr(st, "secrets") else {}
    return {
        "host": secrets_db.get("host", "45.6.155.12"),
        "port": str(secrets_db.get("port", "37000")),
        "database": secrets_db.get("database", "CLEZM5_199407_PR_PD"),
        "user": secrets_db.get("user", "CLT199407readPrime"),
        "password": secrets_db.get("password", ""),
        "driver": secrets_db.get("driver", "ODBC Driver 17 for SQL Server")
    }

def get_connection(cfg: dict):
    conn_str = (
        f"DRIVER={{{cfg['driver']}}};"
        f"SERVER={cfg['host']},{cfg['port']};"
        f"DATABASE={cfg['database']};"
        f"UID={cfg['user']};PWD={cfg['password']};"
        f"TrustServerCertificate=yes;Connection Timeout=20;"
    )
    return pyodbc.connect(conn_str, autocommit=True)

# ─────────────────────────────────────────────
#  Carregamento de Filtros Dinâmicos (Cacheado)
# ─────────────────────────────────────────────
@st.cache_data(ttl=600, show_spinner=False)
def load_filter_options(cfg: dict):
    conn = get_connection(cfg)
    cur = conn.cursor()

    cur.execute("""
        SELECT DISTINCT SUBSTRING(CT2_DATA, 1, 4) AS ANO
        FROM CT2010 (NOLOCK)
        WHERE D_E_L_E_T_ = '' AND LEN(LTRIM(CT2_DATA)) >= 4
          AND SUBSTRING(CT2_DATA, 1, 4) >= '2024'
        ORDER BY ANO DESC
    """)
    anos = [r[0].strip() for r in cur.fetchall() if r[0]]

    cur.execute("SELECT DISTINCT ZA_NOMUSER FROM SZA010 (NOLOCK) WHERE ZA_FILIAL='' AND D_E_L_E_T_='' AND ZA_NOMUSER!='' AND ZA_NOMUSER NOT LIKE '%Erika Maganha%' ORDER BY ZA_NOMUSER")
    resps = [r[0].strip() for r in cur.fetchall() if r[0]]

    cur.execute("SELECT DISTINCT ZA_TIPO FROM SZA010 (NOLOCK) WHERE ZA_FILIAL='' AND D_E_L_E_T_='' AND ZA_TIPO!='' ORDER BY ZA_TIPO")
    za_tipos = [r[0].strip() for r in cur.fetchall() if r[0]]

    cur.execute("SELECT DISTINCT ZA_TIPOCOM FROM SZA010 (NOLOCK) WHERE ZA_FILIAL='' AND D_E_L_E_T_='' AND ZA_TIPOCOM!='' ORDER BY ZA_TIPOCOM")
    za_tipocoms = [r[0].strip() for r in cur.fetchall() if r[0]]

    cur.execute("SELECT CTT_CUSTO, MAX(CTT_DESC01) FROM CTT010 (NOLOCK) WHERE D_E_L_E_T_='' AND CTT_CUSTO!='' AND CTT_BLOQ != '1' GROUP BY CTT_CUSTO ORDER BY CTT_CUSTO")
    ccs = [{"code": r[0].strip(), "label": f"{r[0].strip()} – {_s(r[1])}"} for r in cur.fetchall() if r[0]]

    cur.execute("SELECT CT1_CONTA, MAX(CT1_DESC01) FROM CT1010 (NOLOCK) WHERE D_E_L_E_T_='' AND CT1_CONTA!='' AND SUBSTRING(CT1_CONTA,1,1)>='3' GROUP BY CT1_CONTA ORDER BY CT1_CONTA")
    contas = [{"code": r[0].strip(), "label": f"{r[0].strip()} – {_s(r[1])}"} for r in cur.fetchall() if r[0]]

    conn.close()
    return {
        "anos": anos or ["2026", "2025", "2024"],
        "responsaveis": resps,
        "za_tipos": za_tipos,
        "za_tipocoms": za_tipocoms,
        "ccs": ccs,
        "contas": contas
    }

# ─────────────────────────────────────────────
#  Query SQL Principal
# ─────────────────────────────────────────────
_SQL_BASE = """
SELECT
    CT2.R_E_C_N_O_ AS ROW_ID,
    CT2.CT2_DATA, CT2.CT2_LOTE, CT2.CT2_HIST,
    CT2.CT2_DEBITO,    CT1_D.CT1_DESC01 AS CT1_DESC_DEB,
    CTS_D.CTS_ORDEM AS CTS_ORDEM_DEB, CTS_D.CTS_DESCCG AS CTS_DESCCG_DEB,
    CT2.CT2_CCD,       CTT_D.CTT_DESC01 AS CTT_DESC_DEB,
    SZA_D.ZA_NOMUSER  AS SZA_RESP_DEB,   SZA_D.ZA_TIPO AS SZA_TIPO_DEB, SZA_D.ZA_TIPOCOM AS SZA_TIPOCOM_DEB, CT2.CT2_ITEMD,
    CT2.CT2_CREDIT,    CT1_C.CT1_DESC01 AS CT1_DESC_CRED,
    CTS_C.CTS_ORDEM AS CTS_ORDEM_CRED, CTS_C.CTS_DESCCG AS CTS_DESCCG_CRED,
    CT2.CT2_CCC,       CTT_C.CTT_DESC01 AS CTT_DESC_CRED,
    SZA_C.ZA_NOMUSER  AS SZA_RESP_CRED,  SZA_C.ZA_TIPO AS SZA_TIPO_CRED, SZA_C.ZA_TIPOCOM AS SZA_TIPOCOM_CRED,
    CASE WHEN ZZ4.ZZ4_ANOREF IS NULL OR LTRIM(RTRIM(ZZ4.ZZ4_ANOREF)) = '' THEN 'ORC' + SUBSTRING(CT2.CT2_DATA, 1, 4) ELSE ZZ4.ZZ4_ANOREF END AS ZZ4_ANOREF, ZZ4.ZZ4_CTADES AS ZZ4_DESC_CONTA,
    ZZ4.ZZ4_MES01 AS ORC_JAN, ZZ4.ZZ4_MES02 AS ORC_FEV,
    ZZ4.ZZ4_MES03 AS ORC_MAR, ZZ4.ZZ4_MES04 AS ORC_ABR,
    ZZ4.ZZ4_MES05 AS ORC_MAI, ZZ4.ZZ4_MES06 AS ORC_JUN,
    ZZ4.ZZ4_MES07 AS ORC_JUL, ZZ4.ZZ4_MES08 AS ORC_AGO,
    ZZ4.ZZ4_MES09 AS ORC_SET, ZZ4.ZZ4_MES10 AS ORC_OUT,
    ZZ4.ZZ4_MES11 AS ORC_NOV, ZZ4.ZZ4_MES12 AS ORC_DEZ,
    (ISNULL(ZZ4.ZZ4_MES01,0)+ISNULL(ZZ4.ZZ4_MES02,0)+ISNULL(ZZ4.ZZ4_MES03,0)+
     ISNULL(ZZ4.ZZ4_MES04,0)+ISNULL(ZZ4.ZZ4_MES05,0)+ISNULL(ZZ4.ZZ4_MES06,0)+
     ISNULL(ZZ4.ZZ4_MES07,0)+ISNULL(ZZ4.ZZ4_MES08,0)+ISNULL(ZZ4.ZZ4_MES09,0)+
     ISNULL(ZZ4.ZZ4_MES10,0)+ISNULL(ZZ4.ZZ4_MES11,0)+ISNULL(ZZ4.ZZ4_MES12,0)
    ) AS TOTAL_ORC_ANO,
    CASE 
        WHEN CT2.CT2_CREDIT != '' AND SUBSTRING(CT2.CT2_CREDIT, 1, 1) >= '3' THEN CT2.CT2_VALOR
        WHEN CT2.CT2_DEBITO != '' AND SUBSTRING(CT2.CT2_DEBITO, 1, 1) >= '3' THEN CT2.CT2_VALOR * -1
        ELSE CT2.CT2_VALOR
    END AS VALOR_REALIZADO
FROM CT2010 CT2 (NOLOCK)
LEFT JOIN (SELECT CT1_CONTA, MAX(CT1_DESC01) AS CT1_DESC01 FROM CT1010 (NOLOCK) WHERE D_E_L_E_T_='' GROUP BY CT1_CONTA) CT1_D ON CT1_D.CT1_CONTA=CT2.CT2_DEBITO
LEFT JOIN (SELECT CT1_CONTA, MAX(CT1_DESC01) AS CT1_DESC01 FROM CT1010 (NOLOCK) WHERE D_E_L_E_T_='' GROUP BY CT1_CONTA) CT1_C ON CT1_C.CT1_CONTA=CT2.CT2_CREDIT
OUTER APPLY (
    SELECT TOP 1 CTS_ORDEM, CTS_DESCCG
    FROM CTS010 (NOLOCK)
    WHERE D_E_L_E_T_ = ''
      AND CTS_CODPLA = 'DR1'
      AND CT2.CT2_DEBITO != ''
      AND CT2.CT2_DEBITO BETWEEN CTS_CT1INI AND (CASE WHEN CTS_CT1FIM = '' OR CTS_CT1FIM IS NULL THEN CTS_CT1INI ELSE CTS_CT1FIM END)
    ORDER BY CTS_ORDEM DESC
) CTS_D
OUTER APPLY (
    SELECT TOP 1 CTS_ORDEM, CTS_DESCCG
    FROM CTS010 (NOLOCK)
    WHERE D_E_L_E_T_ = ''
      AND CTS_CODPLA = 'DR1'
      AND CT2.CT2_CREDIT != ''
      AND CT2.CT2_CREDIT BETWEEN CTS_CT1INI AND (CASE WHEN CTS_CT1FIM = '' OR CTS_CT1FIM IS NULL THEN CTS_CT1INI ELSE CTS_CT1FIM END)
    ORDER BY CTS_ORDEM DESC
) CTS_C
LEFT JOIN (SELECT CTT_CUSTO, MAX(CTT_DESC01) AS CTT_DESC01, MAX(CTT_BLOQ) AS CTT_BLOQ FROM CTT010 (NOLOCK) WHERE D_E_L_E_T_='' GROUP BY CTT_CUSTO) CTT_D ON CTT_D.CTT_CUSTO=CT2.CT2_CCD
LEFT JOIN (SELECT CTT_CUSTO, MAX(CTT_DESC01) AS CTT_DESC01, MAX(CTT_BLOQ) AS CTT_BLOQ FROM CTT010 (NOLOCK) WHERE D_E_L_E_T_='' GROUP BY CTT_CUSTO) CTT_C ON CTT_C.CTT_CUSTO=CT2.CT2_CCC
LEFT JOIN SZA010 SZA_D (NOLOCK) ON SZA_D.ZA_FILIAL='' AND SZA_D.ZA_CCUSTO=CT2.CT2_CCD   AND SZA_D.D_E_L_E_T_=''
LEFT JOIN SZA010 SZA_C (NOLOCK) ON SZA_C.ZA_FILIAL='' AND SZA_C.ZA_CCUSTO=CT2.CT2_CCC   AND SZA_C.D_E_L_E_T_=''
LEFT JOIN ZZ4010 ZZ4   (NOLOCK) ON ZZ4.ZZ4_FILIAL='' 
    AND ZZ4.ZZ4_CTACOD = CASE WHEN SUBSTRING(CT2.CT2_DEBITO, 1, 1) >= '3' THEN CT2.CT2_DEBITO ELSE CT2.CT2_CREDIT END
    AND ZZ4.ZZ4_CCUSTO = CASE WHEN SUBSTRING(CT2.CT2_DEBITO, 1, 1) >= '3' THEN CT2.CT2_CCD ELSE CT2.CT2_CCC END
    AND ZZ4.ZZ4_ANOREF = 'ORC' + SUBSTRING(CT2.CT2_DATA, 1, 4) AND ZZ4.D_E_L_E_T_=''
LEFT JOIN CTD010 CTD_D (NOLOCK) ON CTD_D.CTD_FILIAL='' AND CTD_D.CTD_ITEM=CT2.CT2_ITEMD  AND CTD_D.D_E_L_E_T_=''
LEFT JOIN CTD010 CTD_C (NOLOCK) ON CTD_C.CTD_FILIAL='' AND CTD_C.CTD_ITEM=CT2.CT2_ITEMC  AND CTD_C.D_E_L_E_T_=''
LEFT JOIN CTH010 CTH_D (NOLOCK) ON CTH_D.CTH_FILIAL='' AND CTH_D.CTH_CLVL=CT2.CT2_CLVLDB AND CTH_D.D_E_L_E_T_=''
LEFT JOIN CTH010 CTH_C (NOLOCK) ON CTH_C.CTH_FILIAL='' AND CTH_C.CTH_CLVL=CT2.CT2_CLVLCR AND CTH_C.D_E_L_E_T_=''
WHERE CT2.D_E_L_E_T_=''
  AND {date_condition}
  AND (CT2.CT2_CCD != '' OR CT2.CT2_CCC != '')
  AND ISNULL(CTT_D.CTT_BLOQ, '2') != '1'
  AND ISNULL(CTT_C.CTT_BLOQ, '2') != '1'
  AND ({resp_filter})
  AND (
      (CT2.CT2_DEBITO != '' AND SUBSTRING(CT2.CT2_DEBITO, 1, 1) >= '1')
      OR
      (CT2.CT2_CREDIT != '' AND SUBSTRING(CT2.CT2_CREDIT, 1, 1) >= '1')
  )
  {cc_filter}
  {conta_filter}
  {zatipo_filter}
  {zatipocom_filter}
ORDER BY CT2.CT2_DATA
"""

def fetch_data(cfg, anos_list, meses_list, resps, za_tipos, za_tipocoms, ccs, contas, search_str):
    if not anos_list: anos_list = ["2026"]
    
    if meses_list:
        mes_clauses = []
        date_params = []
        for ano in anos_list:
            for mes in meses_list:
                m_num = int(mes)
                last_day = calendar.monthrange(int(ano), m_num)[1]
                mes_clauses.append("CT2.CT2_DATA BETWEEN ? AND ?")
                date_params += [f"{ano}{mes}01", f"{ano}{mes}{last_day:02d}"]
        date_condition = "(" + " OR ".join(mes_clauses) + ")"
    else:
        date_condition = "CT2.CT2_DATA BETWEEN ? AND ?"
        date_params = [f"{min(anos_list)}0101", f"{max(anos_list)}1231"]

    if resps:
        resp_filter = " OR ".join(f"(SZA_D.ZA_NOMUSER LIKE ? OR SZA_C.ZA_NOMUSER LIKE ?)" for _ in resps)
        resp_params = []
        for r in resps: resp_params.extend([f"%{r}%", f"%{r}%"])
    else:
        resp_filter = "1=1"; resp_params = []

    if za_tipos:
        zatipo_filter = f"AND (SZA_D.ZA_TIPO IN ({','.join('?'*len(za_tipos))}) OR SZA_C.ZA_TIPO IN ({','.join('?'*len(za_tipos))}))"
        zatipo_params = list(za_tipos) * 2
    else:
        zatipo_filter = ""; zatipo_params = []

    if za_tipocoms:
        zatipocom_filter = f"AND (SZA_D.ZA_TIPOCOM IN ({','.join('?'*len(za_tipocoms))}) OR SZA_C.ZA_TIPOCOM IN ({','.join('?'*len(za_tipocoms))}))"
        zatipocom_params = list(za_tipocoms) * 2
    else:
        zatipocom_filter = ""; zatipocom_params = []

    if ccs:
        cc_filter = f"AND (CT2.CT2_CCD IN ({','.join('?'*len(ccs))}) OR CT2.CT2_CCC IN ({','.join('?'*len(ccs))}))"
        cc_params = list(ccs) * 2
    else:
        cc_filter = ""; cc_params = []

    if contas:
        conta_filter = f"AND (CT2.CT2_DEBITO IN ({','.join('?'*len(contas))}) OR CT2.CT2_CREDIT IN ({','.join('?'*len(contas))}))"
        conta_params = list(contas) * 2
    else:
        conta_filter = ""; conta_params = []

    sql = _SQL_BASE.format(
        date_condition=date_condition,
        resp_filter=resp_filter,
        cc_filter=cc_filter,
        conta_filter=conta_filter,
        zatipo_filter=zatipo_filter,
        zatipocom_filter=zatipocom_filter
    )

    all_params = date_params + resp_params + zatipo_params + zatipocom_params + cc_params + conta_params
    
    conn = get_connection(cfg)
    cur = conn.cursor()
    cur.execute(sql, all_params)
    cols = [d[0] for d in cur.description]
    all_rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    conn.close()

    if search_str:
        s_low = search_str.lower()
        all_rows = [r for r in all_rows if s_low in " ".join([
            _s(r.get("CT2_HIST")), _s(r.get("CT1_DESC_DEB")), _s(r.get("CT1_DESC_CRED")),
            _s(r.get("CTT_DESC_DEB")), _s(r.get("CTT_DESC_CRED")), _s(r.get("CT2_DEBITO")),
            _s(r.get("CT2_CREDIT")), _s(r.get("CT2_CCD")), _s(r.get("CT2_CCC")),
            _s(r.get("SZA_RESP_DEB")), _s(r.get("SZA_RESP_CRED"))
        ]).lower()]

    gl_rows = []
    for d in all_rows:
        val = _f(d.get("VALOR_REALIZADO"))
        deb_acc = _s(d.get("CT2_DEBITO"))
        deb_cc = _s(d.get("CT2_CCD"))
        if deb_acc and not (deb_acc.startswith('1') or deb_acc.startswith('2')):
            keep = True
            if contas and deb_acc not in contas: keep = False
            if ccs and deb_cc not in ccs: keep = False
            if keep:
                dd = d.copy()
                dd["_is_cred"] = False
                dd["VALOR_REALIZADO"] = -abs(val)
                gl_rows.append(dd)

        cred_acc = _s(d.get("CT2_CREDIT"))
        cred_cc = _s(d.get("CT2_CCC"))
        if cred_acc and not (cred_acc.startswith('1') or cred_acc.startswith('2')):
            keep = True
            if contas and cred_acc not in contas: keep = False
            if ccs and cred_cc not in ccs: keep = False
            if keep:
                dc = d.copy()
                dc["_is_cred"] = True
                dc["VALOR_REALIZADO"] = abs(val)
                gl_rows.append(dc)

    return gl_rows

# ─────────────────────────────────────────────
#  Geração de Excel em Memória
# ─────────────────────────────────────────────
def generate_excel_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Orçado vs Realizado"
    hf = PatternFill("solid", fgColor="1E3A5F")
    hft = Font(bold=True, color="FFFFFF", size=10)
    ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
    bd = Border(left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
                top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))

    headers = ["Data", "Lote", "Histórico", "Conta Déb.", "Desc. Déb.", "CC Déb.", "Desc. CC",
               "Resp. Déb.", "Resp. Créd.", "Conta Créd.", "CC Créd.",
               "Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez",
               "Total Orc. Ano", "Valor Realizado"]
    col_keys = ["CT2_DATA", "CT2_LOTE", "CT2_HIST", "CT2_DEBITO", "CT1_DESC_DEB", "CT2_CCD",
                "CTT_DESC_DEB", "SZA_RESP_DEB", "SZA_RESP_CRED", "CT2_CREDIT", "CT2_CCC",
                "ORC_JAN", "ORC_FEV", "ORC_MAR", "ORC_ABR", "ORC_MAI", "ORC_JUN",
                "ORC_JUL", "ORC_AGO", "ORC_SET", "ORC_OUT", "ORC_NOV", "ORC_DEZ",
                "TOTAL_ORC_ANO", "VALOR_REALIZADO"]
    orc_cols = set(range(12, 25))

    ws.row_dimensions[1].height = 30
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h)
        c.fill = hf; c.font = hft; c.alignment = ha; c.border = bd

    wf = PatternFill("solid", fgColor="FFFFFF")
    af = PatternFill("solid", fgColor="F4F7FB")
    of = PatternFill("solid", fgColor="E8F4FD")
    rf = PatternFill("solid", fgColor="E8F8F0")

    for ri, d in enumerate(rows, 2):
        row_fill = wf if ri % 2 == 0 else af
        for ci, key in enumerate(col_keys, 1):
            val = _s(d.get(key, ""))
            c = ws.cell(ri, ci, val); c.border = bd
            c.alignment = Alignment(vertical="center")
            if ci in orc_cols: c.fill = of
            elif ci == 25: c.fill = rf
            else: c.fill = row_fill
            if ci >= 12:
                try:
                    c.value = float(val)
                    c.number_format = '#,##0.00'
                except: pass

    widths = [12, 10, 45, 14, 28, 10, 28, 24, 24, 14, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 10, 16, 16]
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# ─────────────────────────────────────────────
#  Sidebar e Filtros
# ─────────────────────────────────────────────
st.sidebar.title("🏢 Base Gestor")
st.sidebar.markdown("**Filtros Orçamentários TOTVS Protheus**")

db_cfg = get_db_credentials()

with st.sidebar.expander("⚙️ Conexão com Banco de Dados"):
    cfg_host = st.text_input("Servidor / Host", value=db_cfg["host"])
    cfg_port = st.text_input("Porta", value=db_cfg["port"])
    cfg_db   = st.text_input("Database", value=db_cfg["database"])
    cfg_user = st.text_input("Usuário", value=db_cfg["user"])
    cfg_pwd  = st.text_input("Senha", value=db_cfg["password"], type="password")
    active_cfg = {
        "host": cfg_host, "port": cfg_port, "database": cfg_db,
        "user": cfg_user, "password": cfg_pwd, "driver": db_cfg["driver"]
    }

try:
    filter_opts = load_filter_options(active_cfg)
except Exception as e:
    st.sidebar.error(f"Erro ao conectar no banco: {e}")
    filter_opts = {"anos": ["2026", "2025"], "responsaveis": [], "za_tipos": [], "za_tipocoms": [], "ccs": [], "contas": []}

sel_anos = st.sidebar.multiselect("Ano(s)", options=filter_opts["anos"], default=[filter_opts["anos"][0]] if filter_opts["anos"] else ["2026"])

meses_labels = {f"{m:02d}": f"{m:02d} - {MESES_PT[m]}" for m in range(1, 13)}
sel_meses = st.sidebar.multiselect("Mês(es)", options=list(meses_labels.keys()), format_func=lambda x: meses_labels[x])

sel_resps = st.sidebar.multiselect("Responsável", options=filter_opts["responsaveis"])
sel_za_tipos = st.sidebar.multiselect("Tipo (SZA)", options=filter_opts["za_tipos"])
sel_za_tipocoms = st.sidebar.multiselect("Tipo Comercial", options=filter_opts["za_tipocoms"])

sel_ccs_raw = st.sidebar.multiselect("Centro de Custo", options=filter_opts["ccs"], format_func=lambda x: x["label"])
sel_ccs = [x["code"] for x in sel_ccs_raw]

sel_contas_raw = st.sidebar.multiselect("Conta Contábil", options=filter_opts["contas"], format_func=lambda x: x["label"])
sel_contas = [x["code"] for x in sel_contas_raw]

search_term = st.sidebar.text_input("🔍 Busca rápida (histórico, contas...)")
group_by_dre = st.sidebar.checkbox("Agrupar por Estrutura DRE", value=False)

btn_filtrar = st.sidebar.button("Aplicar Filtros / Buscar", type="primary", use_container_width=True)

# ─────────────────────────────────────────────
#  Processamento Principal
# ─────────────────────────────────────────────
if "data_loaded" not in st.session_state:
    st.session_state.data_loaded = False
    st.session_state.rows = []

if btn_filtrar or not st.session_state.data_loaded:
    with st.spinner("Consultando dados contábeis no SQL Server..."):
        try:
            st.session_state.rows = fetch_data(
                active_cfg, sel_anos, sel_meses, sel_resps,
                sel_za_tipos, sel_za_tipocoms, sel_ccs, sel_contas, search_term
            )
            st.session_state.data_loaded = True
        except Exception as e:
            st.error(f"Erro na consulta SQL: {e}")
            st.session_state.rows = []

rows = st.session_state.rows

def get_row_budget(d):
    if sel_meses:
        return sum(_f(d.get(ORC_KEYS[int(m)-1])) for m in sel_meses)
    return _f(d.get("TOTAL_ORC_ANO"))

def get_grouping_keys(d, use_dre=None):
    if use_dre is None: use_dre = group_by_dre
    is_cred = d.get("_is_cred", False)
    if is_cred:
        k, desc, cc, desc_cc = _s(d.get("CT2_CREDIT")), _s(d.get("CT1_DESC_CRED")), _s(d.get("CT2_CCC")), _s(d.get("CTT_DESC_CRED"))
        dre_ordem, dre_desc = _s(d.get("CTS_ORDEM_CRED")), _s(d.get("CTS_DESCCG_CRED"))
    else:
        k, desc, cc, desc_cc = _s(d.get("CT2_DEBITO")), _s(d.get("CT1_DESC_DEB")), _s(d.get("CT2_CCD")), _s(d.get("CTT_DESC_DEB"))
        dre_ordem, dre_desc = _s(d.get("CTS_ORDEM_DEB")), _s(d.get("CTS_DESCCG_DEB"))
    
    dre_info = f"{dre_ordem} - {dre_desc}".strip(" -") if (dre_ordem or dre_desc) else ""
    if use_dre:
        return (dre_ordem, dre_desc, cc, desc_cc, is_cred) if dre_info else (k, desc, cc, desc_cc, is_cred)
    else:
        if dre_info: desc = f"{desc} (DRE: {dre_info})" if desc else f"DRE: {dre_info}"
        return k, desc, cc, desc_cc, is_cred

# ─────────────────────────────────────────────
#  Cálculo dos Totais e Métricas (Cards)
# ─────────────────────────────────────────────
grp_dict = {}
grp_budget_added = set()
for d in rows:
    k, desc, cc, desc_cc, _ = get_grouping_keys(d)
    if not k: continue
    if k not in grp_dict:
        grp_dict[k] = {"conta": k, "desc": desc, "sum_real": 0.0, "orc": 0.0, "rows": []}
    grp_dict[k]["sum_real"] += _f(d.get("VALOR_REALIZADO"))
    grp_dict[k]["rows"].append(d)

    real_k, _, _, _, _ = get_grouping_keys(d, use_dre=False)
    ano_ref_d = _s(d.get("ZZ4_ANOREF"))
    bkey = (real_k, cc, ano_ref_d)
    if bkey not in grp_budget_added:
        grp_budget_added.add(bkey)
        grp_dict[k]["orc"] += get_row_budget(d)

total_real = sum(g["sum_real"] for g in grp_dict.values())
total_orc  = sum(g["orc"] for g in grp_dict.values())
variacao   = abs(total_real) - total_orc
pct_var    = (variacao / total_orc * 100) if total_orc else 0.0

st.title("📊 Painel Orçamentário - Base Gestor")
col_c1, col_c2, col_c3, col_c4, col_c5 = st.columns(5)
col_c1.metric("Total Lançamentos", f"{len(rows):,}")
col_c2.metric("Total Realizado", _brl(total_real))
col_c3.metric("Total Orçado", _brl(total_orc))
col_c4.metric("Variação (R$)", _brl(variacao), delta=f"{_brl(variacao)}", delta_color="inverse")
col_c5.metric("Variação (%)", _pct(pct_var), delta=f"{_pct(pct_var)}", delta_color="inverse")

st.markdown("---")

tab_detalhe, tab_grade, tab_export = st.tabs(["📑 Lançamentos & Grupos", "📅 Quadro Mensal", "📥 Exportação"])

with tab_detalhe:
    st.subheader("Análise por Grupo Contábil / DRE")
    for k in sorted(grp_dict.keys()):
        g = grp_dict[k]
        g_real = g["sum_real"]
        g_orc = g["orc"]
        g_var = abs(g_real) - g_orc
        g_pct = (g_var / g_orc * 100) if g_orc else 0.0

        header_label = f"**{k}** - {g['desc']} | Real: {_brl(g_real)} | Orç: {_brl(g_orc)} | Var: {_brl(g_var)} ({_pct(g_pct)})"
        with st.expander(header_label, expanded=False):
            sub_rows = []
            for item in g["rows"]:
                sub_rows.append({
                    "Data": _date(item.get("CT2_DATA")),
                    "Lote": _s(item.get("CT2_LOTE")),
                    "Histórico": _s(item.get("CT2_HIST")),
                    "CC Déb": _s(item.get("CT2_CCD")),
                    "Conta Déb": _s(item.get("CT2_DEBITO")),
                    "CC Créd": _s(item.get("CT2_CCC")),
                    "Conta Créd": _s(item.get("CT2_CREDIT")),
                    "Realizado": _f(item.get("VALOR_REALIZADO")),
                })
            df_sub = pd.DataFrame(sub_rows)
            st.dataframe(
                df_sub.style.format({"Realizado": lambda x: _brl(x)}),
                use_container_width=True,
                height=250
            )

with tab_grade:
    st.subheader("Quadro Orçamentário por Mês")
    grid_data = []
    for k in sorted(grp_dict.keys()):
        g = grp_dict[k]
        sample = g["rows"][0] if g["rows"] else {}
        row_dict = {"Conta/Grupo": k, "Descrição": g["desc"], "Realizado": g["sum_real"]}
        for idx, key in enumerate(ORC_KEYS, 1):
            row_dict[MESES_SHORT[idx]] = _f(sample.get(key))
        row_dict["Total Orçado"] = g["orc"]
        grid_data.append(row_dict)
    
    if grid_data:
        df_grid = pd.DataFrame(grid_data)
        st.dataframe(
            df_grid.style.format({c: lambda x: _brl(x) for c in df_grid.columns if c not in ["Conta/Grupo", "Descrição"]}),
            use_container_width=True
        )

with tab_export:
    st.subheader("Exportar Relatório Excel")
    if rows:
        excel_data = generate_excel_bytes(rows)
        st.download_button(
            label="💾 Baixar Relatório Completo (.xlsx)",
            data=excel_data,
            file_name=f"BaseGestor_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary"
        )
    else:
        st.info("Nenhum dado retornado para exportação.")
